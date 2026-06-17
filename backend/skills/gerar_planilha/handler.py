"""Skill gerar_planilha: gera .xlsx (openpyxl) ou .csv a partir de dados
estruturados, sobe ao Storage e devolve um link de download assinado.

As funcoes `_build_*_bytes` sao PURAS (sem dependencia do Storage nem dos tipos
do agente) para permitir teste isolado no host; `openpyxl` e importado de forma
preguicosa (so quando gera .xlsx), entao o modulo carrega so com stdlib.
"""
import csv as _csv
import io


def _build_csv_bytes(colunas: list, linhas: list) -> bytes:
    buf = io.StringIO()
    writer = _csv.writer(buf)
    writer.writerow(list(colunas))
    for linha in linhas:
        writer.writerow(["" if v is None else str(v) for v in linha])
    # utf-8-sig (BOM) para o Excel abrir acentuacao corretamente
    return buf.getvalue().encode("utf-8-sig")


def _build_xlsx_bytes(titulo: str, colunas: list, linhas: list) -> bytes:
    from openpyxl import Workbook  # import preguicoso (dependencia pesada)
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = ((titulo or "Planilha").strip()[:31]) or "Planilha"  # Excel: max 31 chars
    ws.append(list(colunas))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for linha in linhas:
        ws.append(["" if v is None else str(v) for v in linha])
    # largura aproximada das colunas pelo maior conteudo (limitada a 60)
    for i, col in enumerate(colunas, start=1):
        celulas = [len(str(col))] + [len(str(l[i - 1])) for l in linhas if i - 1 < len(l)]
        ws.column_dimensions[get_column_letter(i)].width = min(max(celulas + [8]) + 2, 60)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def executar(args: dict, ctx) -> "object":
    from src.chatbot_fai_docs.agent.types import SkillResult

    titulo = (args.get("titulo") or "").strip()
    colunas = args.get("colunas") or []
    linhas = args.get("linhas") or []
    formato = (args.get("formato") or "xlsx").strip().lower()
    if formato not in ("xlsx", "csv"):
        formato = "xlsx"

    if not titulo or not colunas:
        return SkillResult(
            for_model="Erro: gerar_planilha requer 'titulo' e ao menos uma coluna em 'colunas'.",
            error=True,
        )
    if ctx.storage is None:
        return SkillResult(
            for_model="Erro: armazenamento (Storage) indisponivel para entregar a planilha.",
            error=True,
        )

    try:
        if formato == "csv":
            conteudo = _build_csv_bytes(colunas, linhas)
            ext, content_type = ".csv", "text/csv; charset=utf-8"
        else:
            conteudo = _build_xlsx_bytes(titulo, colunas, linhas)
            ext = ".xlsx"
            content_type = (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
    except Exception as e:  # noqa: BLE001 - erro de geracao vira mensagem p/ o modelo
        return SkillResult(for_model=f"Erro ao montar a planilha: {e}", error=True)

    nome = ctx.storage.sanitize_object_name(titulo) + ext
    try:
        ctx.storage.upload(nome, conteudo, content_type=content_type)
        url = ctx.storage.create_signed_url(nome, expires_in=ctx.signed_url_ttl)
    except Exception as e:  # noqa: BLE001
        return SkillResult(for_model=f"Erro ao salvar/entregar a planilha: {e}", error=True)

    return SkillResult(
        for_model=(
            f"Planilha '{nome}' gerada com {len(linhas)} linha(s) e {len(colunas)} "
            f"coluna(s). O link de download sera anexado automaticamente a resposta."
        ),
        downloads=[(nome, url)],
    )
