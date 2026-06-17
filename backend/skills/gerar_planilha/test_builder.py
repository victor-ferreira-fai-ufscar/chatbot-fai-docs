"""Teste host das funcoes puras de geracao da skill gerar_planilha.

Roda sem Docker:
    python3 backend/skills/gerar_planilha/test_builder.py

Testa CSV sempre (stdlib); testa XLSX apenas se `openpyxl` estiver disponivel
no ambiente. Carrega handler.py por caminho (modulo-level so usa stdlib).
"""
import importlib.util
import io
import sys
from pathlib import Path

_HANDLER = Path(__file__).resolve().parent / "handler.py"
_spec = importlib.util.spec_from_file_location("_gp_handler", _HANDLER)
mod = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(mod)


def main():
    fails = []

    def chk(cond, msg):
        print(("OK   " if cond else "FALHA ") + msg)
        if not cond:
            fails.append(msg)

    colunas = ["Nome", "Valor"]
    linhas = [["Item A", "10"], ["Item B", "20"]]

    # CSV (stdlib, sempre)
    csv_bytes = mod._build_csv_bytes(colunas, linhas)
    text = csv_bytes.decode("utf-8-sig")
    chk(isinstance(csv_bytes, bytes), "csv: retorna bytes")
    chk("Nome,Valor" in text and "Item A,10" in text and "Item B,20" in text,
        "csv: cabecalho e linhas presentes")

    # XLSX (so se openpyxl disponivel)
    try:
        import openpyxl  # noqa: F401
        has_openpyxl = True
    except Exception:
        has_openpyxl = False

    if has_openpyxl:
        xlsx_bytes = mod._build_xlsx_bytes("Relatorio", colunas, linhas)
        chk(isinstance(xlsx_bytes, bytes) and len(xlsx_bytes) > 0, "xlsx: bytes nao-vazios")
        chk(xlsx_bytes[:2] == b"PK", "xlsx: assinatura ZIP/OOXML (PK)")
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        chk([c.value for c in ws[1]] == colunas, "xlsx: cabecalho correto")
        chk(ws.cell(row=2, column=1).value == "Item A", "xlsx: primeira celula de dados")
    else:
        print("SKIP xlsx (openpyxl ausente neste ambiente)")

    print()
    if fails:
        print(f"=== {len(fails)} FALHA(S) ===")
        sys.exit(1)
    print("=== CHECKS OK ===")


if __name__ == "__main__":
    main()
